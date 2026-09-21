from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/Users/mehmetaltay/Documents/ChatGPT/New project/miwae_presentation_optimized/output/Pragmatic_Considerations_MIWAE_Introduction_Reordered_Presenter_Notes_v1_2026-09-21.docx')
LOGO = Path('/Users/mehmetaltay/Documents/ChatGPT/New project/miwae_presentation_optimized/assets/ICLHE_logo_official.jpg')

BLUE = '176082'
DARK_BLUE = '123B53'
ORANGE = 'F0B113'
PALE_BLUE = 'E6F0F5'
LIGHT_GRAY = 'F2F2F5'
BORDER = 'D9D9D9'

slides = [
    (1, 'Pragmatic Considerations for Authors and Publishers on Modification of Input in Written Academic English',
     'Establish the object of inquiry and the status of the evidence.',
     'This presentation asks what authors and publishers can do on the production side of written academic communication. MIWAE is a working framework for examining whether EMI textbooks anticipate multilingual readers while preserving disciplinary precision. The student records and interview quotations shown later are synthetic demonstration data. The literature, located textbook examples, and pilot corpus measurements are source-based.',
     'Do not introduce MIWAE as a validated scale. Call it a provisional coursebook-fit framework and an open analysis prototype.',
     'I will begin with the two traditions that MIWAE brings together: EAP and academic English as a lingua franca.'),
    (2, 'EAP and academic ELF: related fields, different starting points',
     'Define both fields and explain why an ELF perspective became necessary.',
     'Hyland and Shaw define English for Academic Purposes as language research and instruction focused on the communicative needs and practices of people working in academic contexts. It includes academic-text description, literacy, pedagogy, and assessment. Seidlhofer defines English as a lingua franca as English used among speakers of different first languages when English is the communicative medium of choice, often the only option. Academic ELF applies that perspective to higher education and research.',
     'Mauranen, Hynninen, and Ranta identify a historical contrast in norm orientation. EAP traditionally often used native-speaker English as a reference, whereas academic ELF studies international settings in which most participants use English as an additional language. ELFA emerged because successful global academic communication cannot be explained only through closeness to an idealised native norm. Present this as a difference of starting orientation; EAP is diverse, and academic ELF does not reject form or standards.',
     'The clearest historical contrast concerns mode: EAP developed a strong writing tradition, while early academic ELF research began with speech.'),
    (3, 'Writing and speech were starting tendencies—not boundaries',
     'Use Mauranen et al. precisely, then establish the written-ELF gap.',
     'Mauranen and colleagues state: “EAP has a strong tradition of focusing on written language, whereas academic ELF research started from analysing spoken discourse.” Early academic ELF research prioritised speech because researchers wanted to observe how speakers manage communication when English is an additional language for everyone, and because language change is often first visible in interaction.',
     'This is a research-history contrast rather than a permanent division. Björkman notes that present-day ELF is used in both spoken and written form. Mauranen and colleagues describe WrELFA as a corpus created to make written academic ELF observable, but they also report that only a handful of studies had addressed writing and call written academic ELF almost entirely uncharted. Written communication therefore belongs within ELFA, while remaining less studied than speech.',
     'EMI textbooks make that imbalance practically important because they carry disciplinary knowledge through sustained written communication.'),
    (4, 'EMI textbooks make written academic communication central',
     'Present the textbook as a communicative interface and bound the evidence carefully.',
     'The slide models a chain: authors, editors, and publishers choose wording, relations, sequence, examples, visuals, and notation; the textbook stages disciplinary explanations; multilingual students interpret those choices; and disciplinary knowledge is reconstructed through reading. A textbook is therefore both a knowledge resource and a written communicative interface.',
     'Jenkins and Mauranen’s volume reports the use of textbooks published by Anglo-American universities in a Chinese EMI context. In the courses examined, authentic English textbooks were described as the cornerstone, often with Chinese supplementary material. This evidence is context-specific. It supports studying textbooks as a primary route to knowledge in some EMI settings, but it does not show that every EMI programme is textbook-led or that all EMI textbooks are produced in Anglo-American countries.',
     'The production-side question follows: do authors, editors, and publishers calibrate language and design for a multilingual global readership?'),
    (5, 'Can textbook prose be treated as input?',
     'Introduce input modification without importing SLA assumptions uncritically.',
     'Written textbook prose can legitimately be called linguistic input. Second-language reading studies explicitly compare unmodified, simplified, and elaborated written passages. Li, Xu, and Wang’s study is a clear example of research on modified written input and reading comprehension.',
     'MIWAE makes an operational extension. The textbook input carries disciplinary knowledge, so the purpose is communicative access to precise content rather than language acquisition alone. Modification means purposeful calibration: defining or glossing terms, signalling relations, sequencing explanations, adding examples, reactivating prerequisite context, and aligning prose with visuals or notation. It may involve local simplification, elaboration, or redistribution of support. It must not erase technical distinctions or reduce disciplinary standards.',
     'This relational account of reader and text connects to McKinley’s argument that EMI preparedness involves academic literacy as well as general English proficiency.'),
    (6, 'Academic literacy links reader readiness to textbook design',
     'Use McKinley to link learner readiness with the design environment.',
     'McKinley argues that EMI challenges often stem from gaps in academic literacy rather than linguistic deficits alone. Preparedness includes learning to think, read, write, reason, and interpret evidence in discipline-specific ways. Standardised proficiency scores have limited predictive value when used by themselves.',
     'McKinley’s argument concerns institutional and pedagogical support rather than a textbook-rating system. MIWAE makes a careful connection: the textbook is one part of the academic-literacy environment through which disciplinary expectations become visible or remain implicit. A student can have adequate general English but still struggle with unstated rhetorical relations, epistemic expectations, or the coordination of prose with disciplinary representations. Text design cannot compensate for every gap in prior knowledge, but it can support or obstruct disciplinary reading.',
     'The CEFR’s treatment of pragmatic competence helps specify which communicative dimensions of textbook design can be examined.'),
    (7, 'CEFR pragmatic competence offers a three-part design lens',
     'Explain the CEFR components and mark the textbook mapping as MIWAE’s extension.',
     'The CEFR Companion Volume says pragmatic competence concerns actual language use in the co-construction of text. Messages are organised, structured, and arranged through discourse competence; used to perform communicative functions through functional competence; and sequenced according to interactional and transactional schemata through design competence.',
     'The CEFR describes competence of users and learners. It does not provide a textbook scale. MIWAE makes a production-side mapping: discourse motivates questions about cohesion, information structure, and rhetorical relations; functional competence motivates questions about definitions, explanations, contrasts, and reader guidance; and design competence motivates questions about sequencing, placement, distribution, and multimodal coordination. Present this mapping as an explicit analytical inference.',
     'These links now allow the research problem and MIWAE’s bounded contribution to be stated precisely.'),
    (8, 'Research problem and MIWAE contribution',
     'State the question, the evidence gap, and the intended decision product.',
     'The research question is: How can EMI textbooks be examined for the way they modify and organise written disciplinary input for multilingual readers? The problem has three layers: written academic ELF remains under-researched; textbook discourse is rarely examined as written communication with a multilingual audience; and reported reader difficulty is seldom linked to located design choices in the text.',
     'MIWAE proposes to connect reader evidence with observable textual evidence, organise that evidence through function, form, and distribution, and support decisions by authors, publishers, syllabus designers, and course teams. ELFA-friendly is a project term for design that anticipates multilingual readers while preserving disciplinary precision. The framework remains provisional and requires reliability and validity work; it is not a universal ranking of textbooks.',
     'I will now show how the synthetic reader evidence and the ten-book pilot were organised.'),
    (9, 'Demonstration design and corpus scope',
     'Explain the mixed evidence design and sampling logic.',
     'The demonstration contains 100 fictional second-year records, 20 fictional semi-structured interviews, ten reviewed textbooks, and 50 matched passages. The ten-book set is balanced: five Electrical Engineering and five International Relations titles. Nine books have complete searchable text layers; the Mano scan contributes five stratified OCR passages.',
     'Second-year students are defensible because nearly all completed a one-year preparatory English programme and are now in their first sustained encounter with disciplinary EMI reading. This captures the transition from general English to discipline-specific academic literacy, while difficulties and support episodes remain recent. Fourth-year students have more experience, but their coping routines may conceal early design barriers. The design targets the transition into disciplinary reading rather than maximum years of exposure.',
     'The framework-development sequence was: literature-based sensitising functions; neutral open coding of student accounts; textbook coding of forms and bounded distributions; readability and CEFR estimates as an external language lens; and revision into the MIWAE decision tool. Neutral elicitation must precede candidate-feature prompts.',
     'The analysis pipeline is implemented in an open prototype API.'),
    (10, 'MIWAE Corpus Analysis API',
     'Explain what the tool does and where human judgement remains necessary.',
     'The API accepts text, Markdown, PDF, DOCX, or grouped CSV; segments the material; profiles candidate functions and forms; calculates distribution measures; and exports records for contextual review. The pilot included ten books, 50 stratified passages, 9,810 words, and 260 book-feature rows.',
     'Text Inspector exports and EVP or EGP mappings are optional user-supplied inputs. Do not claim an official live connection to these services. Automatic matches are candidates until a researcher checks the local rhetorical function and context.',
     'The next three slides show access, execution, and output.'),
    (11, 'From GitHub to a running local API',
     'Show that the research tool is public and reproducible.',
     'The public repository is https://github.com/MehmetDEU/MIWAE. A user clones the repository, creates a virtual environment, installs the package, starts Uvicorn, and then opens the local Swagger interface at http://localhost:8000/docs. The Terminal process must remain running.',
     'GitHub is the public source location. Localhost is a service running only on the user’s own computer. If time is short, describe the three stages without performing the installation live.',
     'Once Swagger opens, a passage can be submitted in four steps.'),
    (12, 'Four steps turn a passage into an analysis',
     'Demonstrate the interface without showing an AI application.',
     'Open POST /v1/analyze/text, select Try it out, paste the request JSON into the Request body editor, and select Execute. The endpoint label itself is clicked; it is not pasted. Query controls determine whether example hits are returned and how many are shown.',
     'For a live demonstration, use the short circuit passage already shown in the slide. Keep the example short so the audience can see the entire request and response.',
     'The response then separates evidence layers instead of returning a single unexplained score.'),
    (13, 'The response separates six inspectable evidence layers',
     'Explain the output hierarchy and the limits of a short demonstration passage.',
     'The response contains the document record, statistics, functions, grammatical forms, distribution rows, segments, and optional external profiles. In the demonstration, a 44-word passage produces six function hits, ten form hits, and 26 distribution rows. High per-10,000 rates are expected with such a short denominator, so this output illustrates structure rather than book-level prevalence.',
     'The API currently returns JSON. The formatted Excel workbook is a companion conversion with Summary, Functions, Forms, Distribution, Segments, and Hits worksheets; it is not yet a built-in export endpoint. Keep the JSON as the machine-readable research record.',
     'I will now move from infrastructure to the demonstration findings.'),
    (14, 'Reading challenge profile in the synthetic dataset',
     'Illustrate how challenge, proficiency, and GPA can be presented together.',
     'The fictional dataset gives an overall reading-challenge mean of 3.19 on the six-point scale. Specific vocabulary and difficult words are the highest items. Challenge is negatively related to institutional English score and GPA in the generated data.',
     'These values are synthetic and descriptive. Do not report p-values, causal effects, or generalise to the target institution. In the eventual study, report reliability, assumptions, confidence intervals, and the exact institutional score metric.',
     'The interviews show what a scale score alone cannot reveal: where meaning breaks down and what forms of support students use.'),
    (15, 'Engineering readers used equations and visuals as parallel routes to meaning',
     'Illustrate disciplinary meaning making in Engineering.',
     'The fictional excerpts show three functions: multimodal verification, precision through labels, and conceptual bridging between procedural steps. Equations, circuits, graphs, and symbols can provide a second route to meaning, but only when their labels and relations are explicit.',
     'The excerpts are generated demonstration language and must not be attributed to real participants. Preserve the ordinary spoken style because it illustrates how participants might explain a reading episode.',
     'International Relations readers describe a different pattern of support.'),
    (16, 'International Relations readers negotiated competing meanings',
     'Illustrate the role of perspective, qualification, and comparison.',
     'The fictional excerpts highlight term ownership, rhetorical relations, and qualified visual support. Concepts such as power and security can have different meanings across theoretical traditions. A comparison table may orient the reader, but it cannot replace prose that explains overlap, disagreement, and uncertainty.',
     'The broader qualitative synthesis contains four patterns: relations, placement, disciplinary meaning, and mode or discipline. Retain negative cases: some difficulties arise from mathematics, missing prior knowledge, limited time, or rushed reading rather than textbook design.',
     'The contrast should be described through semiotic resources and epistemic organisation.'),
    (17, 'Disciplinary difference changes the form of useful support',
     'Make the disciplinary comparison without using a simple objective-versus-subjective binary.',
     'Engineering often coordinates prose, formulae, diagrams, and worked procedures. International Relations often coordinates theoretical perspective, historical context, and qualified argument. These are tendencies in the sampled materials, not fixed properties of all texts in either field.',
     'Avoid saying that Engineering is purely objective and International Relations merely subjective. The defensible claim is that the fields organise evidence and meaning through different semiotic and epistemic resources. MIWAE therefore evaluates disciplinary fit.',
     'The next slide shows how the seven functions appear in located textbook examples.'),
    (18, 'Located textbook examples across the seven functions',
     'Connect abstract functions to observable forms in the textbook corpus.',
     'Walk down the table quickly: definitions can support meaning access; Reader’s Guides can provide orientation; contrast markers can expose rhetorical relations; familiar labels can supply background knowledge; recall cues can reactivate prior concepts; embedded exercises can support self-monitoring; and figure references with labels can coordinate modes.',
     'Additional located examples moved from the slides include Mingst’s in-text explanation of fracking for background knowledge, Mingst’s “as discussed in Chapter 3” for conceptual continuity, and Baylis’s questions inviting readers to reconsider globalisation for reader engagement. A surface form does not determine its function; local rhetorical purpose must be coded.',
     'Located examples illustrate function and form. Distribution requires systematic counting across eligible opportunities.'),
    (19, 'Visual scaffolding varied sharply by discipline',
     'Show convergence between the synthetic reader account and the corpus proxy.',
     'Eight of ten fictional Engineering interviewees mention a graph, table, diagram, or worked visual as helpful, while none of the ten fictional International Relations interviewees does so spontaneously. In the nine complete text layers, figure-reference density is much higher in the Engineering books.',
     'Figure references are a proxy for visual apparatus, not a direct image count. Mano is excluded from the whole-text calculation because only five OCR passages are available. Treat the convergence as a criterion-generating observation rather than a validated causal finding.',
     'Language-analysis tools describe another part of the picture: conventional prose burden.'),
    (20, 'IR prose carried more conventional readability burden',
     'Present the matched-passage comparison while limiting the interpretation.',
     'Across 25 passages per field, the International Relations sample has higher Flesch–Kincaid grade and Gunning Fog values, lower reading ease, longer sentences, and substantially higher VOCD lexical diversity. Academic-word proportions are slightly higher in Engineering, and metadiscourse means are similar.',
     'Readability and estimated CEFR describe linguistic form. They do not measure conceptual accessibility, visual support, pragmatic competence, or ELFA friendliness. The passage-level design is descriptive and needs independent resampling. CEFR spans report the minimum and maximum labels across five passages; D1 is the platform band above C2. Mano is represented by OCR passages and cannot support whole-book frequency claims.',
     'Book-level detail moved from the slide is provided in the tables immediately after these notes.',
     'The discussion integrates reader evidence and textual evidence without collapsing them.'),
    (21, 'Language burden and support operate through different routes',
     'Integrate the disciplinary findings and state the boundaries of the claim.',
     'Engineering readers valued figures, labels, and bridges between procedural steps; the full-text corpus also contained more figure references. International Relations readers valued perspective marking, qualification, and comparison; its matched prose sample used longer sentences and had higher lexical diversity.',
     'Readability cannot establish ELFA friendliness. Figure references are proxies. Synthetic interview patterns illustrate the framework but cannot validate it. A coursebook decision must combine reader evidence with function, form, placement, distribution, and disciplinary fit.',
     'This integration leads to the framework’s three analytic questions.'),
    (22, 'Function, form and distribution answer different questions',
     'Clarify the three-part analytic structure, especially distribution.',
     'Function asks which reader need is addressed. Form asks which observable textual, visual, or numeric device supplies the support. Distribution asks how reliably the device is available across eligible points of need. Reader evidence asks whether students notice, use, or still need the support.',
     'Keep textual and perceived distribution separate. Textual distribution records coverage, timing, and spread in the sampled book. Perceived distribution records how many participants mention a need or support. Distribution requires an opportunity denominator: for example, four supported first introductions out of six eligible first introductions.',
     'The next slide lists the forms and denominators that make these questions codeable.'),
    (23, 'Forms and distributions use explicit coding indicators',
     'Present the coding inventory as open and conditional.',
     'Each function is linked to observable forms and to a distribution rule. Meaning access uses definitions, glosses, paraphrases, or worked examples. Orientation uses previews, headings, recaps, and cross-references. Rhetorical relations use contrast, cause, qualification, and importance cues. The remaining functions use context notes, analogies, recall cues, questions, feedback, labels, captions, worked graphs, and aligned notation.',
     'The forms are not a closed checklist. An eligible opportunity exists only where the communicative need occurs. Define the denominator before coding and record proximity, timing, recurrence, and spread.',
     'A six-step review procedure turns this inventory into a reproducible book-selection process.'),
    (24, 'MIWAE coursebook-fit review in six steps',
     'Show how an author, publisher, or syllabus designer would use MIWAE.',
     'First profile the course and readers. Second set priorities before opening the books. Third sample the same types of locations in every candidate. Fourth code function, form, and distribution. Fifth rate the support and tag the evidence. Sixth choose, supplement, or reconsider the candidate.',
     'Use identical sampling rules and priorities for all candidates to prevent post hoc scoring. The same procedure can support manuscript review by authors and publishers: define the intended readers, identify essential dimensions, inspect eligible points, and revise gaps before publication.',
     'The rubric makes the decision criteria explicit.'),
    (25, 'Coursebook-fit rubric, dimensions 1–4',
     'Explain the first four dimensions and the meaning of the scale.',
     'The 0–3 ratings distinguish unsupported, limited, adequate, and systematic support. The first four dimensions are meaning access, orientation, rhetorical relations, and background knowledge. A systematic score requires consistent alignment across the sampled locations, not the presence of one good example.',
     'Rate sampled evidence rather than an impression of the entire book. Use not applicable only when the dimension is genuinely irrelevant to the course and sample.',
     'The remaining dimensions cover continuity, engagement, and multimodal support.'),
    (26, 'Coursebook-fit rubric, dimensions 5–7 and gates',
     'Complete the rubric and explain the non-negotiable quality gates.',
     'Conceptual continuity asks whether prerequisites are reactivated when needed. Reader engagement asks whether readers can monitor understanding through relevant questions, checks, or worked reasoning. Multimodal support asks whether prose, visuals, and notation explain one another.',
     'Four gates override a high support score: disciplinary accuracy, preserved precision, absence of a native-speaker deficit assumption, and fit with the course’s epistemic and multimodal practices. An inaccurate simplification cannot receive an acceptable coursebook-fit decision.',
     'The final decision rule separates textual evidence from reader and comprehension evidence.'),
    (27, 'Evidence tags and the coursebook decision rule',
     'Explain how evidence strength and local priorities determine the decision.',
     'T means located textual evidence with a sampled denominator. TP adds participant perception. TPC adds comprehension or outcome evidence. A book is chosen when every essential dimension scores at least two and all gates pass. It is chosen with supplementation when one essential dimension scores one and the gap can be covered explicitly. It is reconsidered when an essential dimension scores zero, placement repeatedly fails, or an accuracy gate fails.',
     'Set E, S, and N/A priorities before review. Compare candidates dimension by dimension. Do not publish a universal total score until the instrument has been validated.',
     'The conclusion returns to the pragmatic contribution for authors and publishers.'),
    (28, 'ELFA friendliness becomes reviewable when reader needs are tied to observable textbook choices',
     'Close with the contribution and invite critique of the framework.',
     'MIWAE positions textbook review between written EAP and communication-oriented academic ELF. The ten-book pilot shows why conventional language burden, visual support, and disciplinary meaning must be interpreted together. The rubric and open API make the review process inspectable while preserving disciplinary precision.',
     'Invite the audience to test the dimensions and decision rule against books they write, publish, select, or teach from. Ask which dimensions are missing and which opportunity denominators are feasible in real editorial workflows.',
     'If time permits, use the final two slides for evidence status and references; otherwise keep them for questions.'),
    (29, 'Evidence status and next validation steps',
     'State the project’s current limits transparently.',
     'The student records and interviews are synthetic. The corpus contains ten books, with nine complete text layers and one OCR sample. The readability comparison uses 50 passages. The rubric is a provisional 0–3 decision aid. The API is an open prototype. Comprehension effects have not yet been tested.',
     'The next steps are real data collection with confirmed ethics and translation procedures, complete OCR and edition checks, independent passage sampling, double coding and inter-rater agreement, parser adjudication and codebook versioning, and comprehension comparisons between original and revised passages.',
     'The presentation demonstrates a research and selection procedure; it does not report completed empirical findings.',
     'The final slide lists the main sources and can remain visible during questions.'),
    (30, 'Selected references',
     'Provide traceability and close the deck cleanly.',
     'The most important sources for the spoken argument are Mauranen, Hynninen, and Ranta for academic ELF; Björkman and Seidlhofer for communicative function and accommodation; Jenkins and Mauranen for EMI coursebook difficulty; McKinley for discipline-specific academic literacy; the CEFR Companion Volume for pragmatic competence; and Biber et al. and Quirk et al. for grammatical form inventories.',
     'Verify full bibliographic details, DOI information, textbook editions, and page locators in the conference paper and handout. The ICLHE logo and palette come from the official association website, https://www.iclhe.org/.',
     'Thank the audience and invite questions about the framework, its denominators, and editorial feasibility.'),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_borders(cell, color=BORDER, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_char1, instr, fld_char2])


def add_lead(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(text)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for style_name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 13)]:
    st = styles[style_name]
    st.font.name = 'Arial'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True

header = sec.header.paragraphs[0]
header.text = 'MIWAE Presenter Notes'
header.style = styles['Normal']
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(98, 103, 119)
page_number(sec.footer.paragraphs[0])

if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(LOGO), width=Inches(1.65))

title = doc.add_paragraph(style='Title')
title.add_run('Presenter Notes for Pragmatic Considerations for Authors and Publishers on Modification of Input in Written Academic English')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = subtitle.add_run('MIWAE conference presentation · 30-slide introduction-reordered working edition')
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(23, 96, 130)

intro = doc.add_paragraph()
intro.add_run('Purpose. ').bold = True
intro.add_run('These notes preserve the theoretical, methodological, and technical detail removed from the slides during shortening. They support a concise oral delivery while keeping the distinctions that protect the argument: academic literacy is broader than general proficiency; CEFR pragmatic competence is mapped to textbook design as a MIWAE inference; and function, form, distribution, and reader evidence remain separate analytic layers.')
schedule = doc.add_paragraph()
schedule.add_run('Thirty-minute session plan. ').bold = True
schedule.add_run('Aim for a 20-minute presentation, reserve 8 minutes for questions, and keep a 2-minute buffer for transitions or technical delay. Slides 29–30 are backup slides and do not need full delivery in the main talk.')

doc.add_page_break()
doc.add_heading('Suggested timing', level=1)
timing = doc.add_table(rows=1, cols=3)
timing.alignment = WD_TABLE_ALIGNMENT.CENTER
timing.autofit = False
widths = [Inches(1.2), Inches(2.2), Inches(3.6)]
headers = ['Time', 'Slides', 'Delivery']
for i, text in enumerate(headers):
    cell = timing.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_shading(cell, DARK_BLUE)
    set_cell_borders(cell)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(text)
    rr.bold = True
    rr.font.color.rgb = RGBColor(255, 255, 255)
for vals in [
    ('7 min', '1–8', 'EAP–ELFA contrast, written-ELF gap, input modification, academic literacy, CEFR, and question'),
    ('3 min', '9–13', 'Demonstration design, API access, Swagger workflow, and response structure'),
    ('5 min', '14–21', 'Synthetic student evidence, textbook examples, corpus comparison, and discussion'),
    ('5 min', '22–28', 'Framework, rubric, decision rule, and conclusion'),
    ('Backup', '29–30', 'Evidence status and references for questions'),
    ('8 min', 'Questions', 'Framework validity, distribution denominators, API use, and editorial feasibility'),
    ('2 min', 'Buffer', 'Room transition, technical delay, or one extended answer'),
]:
    row = timing.add_row()
    for i, text in enumerate(vals):
        cell = row.cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, 'FFFFFF' if len(timing.rows) % 2 else LIGHT_GRAY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].add_run(text)

doc.add_heading('Delivery choices for a shorter slot', level=1)
for text in [
    'For a 15-minute slot, summarise slides 2–3 together, compress slides 5–7 into one conceptual bridge, show only one interview-excerpt slide, and present slides 25–26 as a single rubric spread without reading each cell.',
    'If the API cannot be demonstrated live, show slides 11–13 and state that the repository is public while localhost runs only on the presenter’s computer.',
    'Keep slides 29–30 as backup. Their content is already preserved in these notes.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

nav_note = doc.add_paragraph()
nav_run = nav_note.add_run('Navigation in these notes. ')
nav_run.bold = True
nav_run.font.color.rgb = RGBColor(23, 96, 130)
nav_note.add_run('Every section begins with the PowerPoint slide number. An orange instruction at the end of each section tells you exactly when to advance and names the next slide.')

doc.add_page_break()
doc.add_heading('Slide by slide notes', level=1)

for idx, item in enumerate(slides):
    num, title_text, purpose, wording, caution, *remaining = item
    transition = remaining[-1]
    extra_details = remaining[:-1]
    h = doc.add_heading(f'Slide {num}  {title_text}', level=2)
    if num in {7, 8}:
        h.paragraph_format.page_break_before = True
    h.paragraph_format.keep_with_next = True
    add_lead(doc, 'Purpose.', purpose)
    add_lead(doc, 'Suggested wording.', wording)
    add_lead(doc, 'Evidence and caution.', caution)
    for extra in extra_details:
        add_lead(doc, 'Detail retained from the longer deck.', extra)
    p = add_lead(doc, 'Transition.', transition)
    p.runs[-1].italic = True
    cue = doc.add_paragraph()
    cue.paragraph_format.space_before = Pt(2)
    cue.paragraph_format.space_after = Pt(8)
    if idx + 1 < len(slides):
        next_num = slides[idx + 1][0]
        next_title = slides[idx + 1][1]
        cue_text = f'ADVANCE TO SLIDE {next_num}  {next_title}'
    else:
        cue_text = 'END OF DECK  Leave the references visible for questions or return to Slide 28 for the concluding framework.'
    cue_run = cue.add_run(cue_text)
    cue_run.bold = True
    cue_run.font.size = Pt(9.5)
    cue_run.font.color.rgb = RGBColor(213, 140, 0)

    if num == 20:
        doc.add_page_break()
        doc.add_heading('Book level profiles retained in the notes', level=2)
        profiles = [
            ('Hambley Electrical Engineering', 'EE', 'B2+–D1', '12.67', '63.13'),
            ('Mano Digital Logic', 'EE', 'B2–C1+', '10.13', '53.51'),
            ('Oppenheim Signals and Systems', 'EE', 'C1+–D1', '16.37', '62.86'),
            ('Sadiku Electric Circuits', 'EE', 'B2+–C2+', '13.22', '73.39'),
            ('Sedra Microelectronic Circuits', 'EE', 'C1–C2+', '12.85', '54.70'),
            ('Baylis Globalization of World Politics', 'IR', 'C1+–D1', '15.32', '94.71'),
            ('Blanton World Politics', 'IR', 'C2–D1', '16.45', '112.93'),
            ('Dunne International Relations Theories', 'IR', 'C1+–D1', '15.46', '99.17'),
            ('Goldstein International Relations Brief', 'IR', 'C1–C2+', '12.91', '97.00'),
            ('Mingst Essentials of International Relations', 'IR', 'C1–C2+', '14.75', '95.94'),
        ]
        tb = doc.add_table(rows=1, cols=5)
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        ws = [Inches(3.0), Inches(0.55), Inches(1.0), Inches(0.75), Inches(0.8)]
        for i, text in enumerate(['Book', 'Field', 'Observed CEFR span', 'FK grade', 'VOCD']):
            c = tb.rows[0].cells[i]
            c.width = ws[i]
            set_cell_shading(c, DARK_BLUE)
            set_cell_borders(c)
            set_cell_margins(c, top=70, bottom=70)
            rr = c.paragraphs[0].add_run(text)
            rr.bold = True
            rr.font.color.rgb = RGBColor(255, 255, 255)
            rr.font.size = Pt(8.5)
        for ridx, vals in enumerate(profiles):
            row = tb.add_row()
            for i, text in enumerate(vals):
                c = row.cells[i]
                c.width = ws[i]
                set_cell_shading(c, 'FFFFFF' if ridx % 2 == 0 else LIGHT_GRAY)
                set_cell_borders(c)
                set_cell_margins(c, top=65, bottom=65)
                rr = c.paragraphs[0].add_run(text)
                rr.font.size = Pt(8.2)
                if i > 0:
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if num % 2 == 0 and num != 30:
        doc.add_page_break()
    else:
        doc.add_paragraph()

doc.add_heading('Core references for oral attribution', level=1)
refs = [
    'Björkman, B. (2013). English as an Academic Lingua Franca.',
    'Council of Europe. (2020). Common European Framework of Reference for Languages Companion Volume.',
    'Jenkins, J., and Mauranen, A. (Eds.). (2019). Linguistic Diversity on the EMI Campus.',
    'Mauranen, A., Hynninen, N., and Ranta, E. (2016). English as the academic lingua franca.',
    'McKinley, J. (2025). Beyond Proficiency: Rethinking Preparedness in English-Medium Instruction. TESOL Journal, e70080.',
    'Seidlhofer, B. (2011). Understanding English as a Lingua Franca.',
    'ICLHE Association. Official website and visual identity. https://www.iclhe.org/',
]
for ref in refs:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(ref)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
